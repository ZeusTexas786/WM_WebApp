"""Bounded, text-only extraction. Public deployment still requires OS sandboxing."""
import json
import subprocess
import sys
from pathlib import Path
from zipfile import ZipFile

def extract_text(path, format):
    try:
        if format == '.txt':
            text = Path(path).read_text(encoding='utf-8-sig')
            if '\x00' in text:
                raise ValueError('Binary data is not a text resume.')
            segments = [(f'section {i+1}',part.strip()) for i,part in enumerate(text.split('\n\n')) if part.strip()]
        elif format == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(path, strict=True)
            if reader.is_encrypted or len(reader.pages) > 50:
                raise ValueError('Use an unencrypted resume with at most 50 pages.')
            segments = [(f'page {i+1}', page.extract_text() or '') for i, page in enumerate(reader.pages)]
        elif format == '.docx':
            with ZipFile(path) as archive:
                infos = archive.infolist()
                if len(infos) > 300 or sum(item.file_size for item in infos) > 20*1024*1024:
                    raise ValueError('Document expands beyond the safe extraction limit.')
                if any('vbaProject' in item.filename for item in infos):
                    raise ValueError('Macro content is not supported.')
            from docx import Document
            doc = Document(path)
            paragraphs = [p.text for p in doc.paragraphs]
            paragraphs += [' | '.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
            segments = [(f'paragraph {i+1}', value) for i, value in enumerate(paragraphs)]
        else:
            raise ValueError('Unsupported resume format.')
        blocks = []
        total = 0
        for location, text in segments:
            text = text.strip()
            total += len(text)
            if total > 100000 or len(blocks) >= 500:
                raise ValueError('Resume text exceeds the extraction limit.')
            if text:
                blocks.append({'id': f'b{len(blocks)+1}', 'location': location, 'text': text})
        if total < 15:
            raise ValueError('No readable resume text. Scanned PDFs need OCR, which is not included yet.')
        return blocks
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('The document could not be read. Check the format or try a text-based file.') from exc

def parse_bounded(path, format, timeout=20):
    try:
        result = subprocess.run([sys.executable, '-I', str(Path(__file__).resolve()), path, format], capture_output=True, timeout=timeout, check=False)
    except subprocess.TimeoutExpired as exc:
        raise ValueError('Document extraction timed out.') from exc
    if result.returncode or len(result.stdout) > 1500000:
        raise ValueError('Unable to extract this document safely. It may be corrupt, scanned, or too complex.')
    return json.loads(result.stdout)

def parse_remote(path,format,socket_path,timeout=25):
    import socket,struct
    if format not in ('.txt','.pdf','.docx') or not socket_path:
        raise ValueError('Isolated parser is not configured.')
    try:
        with socket.socket(socket.AF_UNIX,socket.SOCK_STREAM) as client:
            client.settimeout(timeout)
            client.connect(socket_path)
            size=Path(path).stat().st_size
            if not 0<size<=10*1024*1024: raise ValueError('Invalid parser input size.')
            header=json.dumps({'format':format,'size':size}).encode()
            client.sendall(struct.pack('!I',len(header))+header)
            with Path(path).open('rb') as stream:
                while chunk:=stream.read(65536): client.sendall(chunk)
            response=b''
            while True:
                chunk=client.recv(65536)
                if not chunk: break
                response+=chunk
                if len(response)>1500000: raise ValueError('Parser response too large.')
        result=json.loads(response)
        blocks=result.get('blocks')
        if not isinstance(blocks,list) or not 1<=len(blocks)<=500: raise ValueError('Isolated parser rejected this document.')
        if any(not isinstance(b,dict) or not isinstance(b.get('text'),str) or not isinstance(b.get('location'),str) or b.get('id')!=f'b{i+1}' for i,b in enumerate(blocks)): raise ValueError('Invalid isolated parser output.')
        if sum(len(b['text']) for b in blocks)>100000: raise ValueError('Parser output exceeds text limit.')
        return blocks
    except (OSError,json.JSONDecodeError,AttributeError,KeyError) as exc:
        raise ValueError('Isolated document extraction is unavailable. No local fallback is permitted.') from exc

if __name__ == '__main__':
    try:
        print(json.dumps(extract_text(sys.argv[1], sys.argv[2])))
    except Exception:
        sys.exit(2)
